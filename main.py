"""Turn a folder of invoice PDFs into one Excel workbook.

Usage:

    python main.py sample_input
    python main.py sample_input -o output/invoices.xlsx

The workbook has three sheets:
- Invoices: one row per file, with status OK or LỖI / ERROR
- Line Items: one row per item
- Issues: one row per problem, in Vietnamese and English

ERROR rows on the Invoices sheet are filled red. Money columns use the
Excel format #,##0.
"""

import argparse
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from extractor import extract_invoice
from validator import validate_invoice

INVOICE_COLUMNS = [
    "File",
    "Mẫu số",
    "Ký hiệu",
    "Số HĐ",
    "Ngày",
    "Người bán",
    "MST người bán",
    "Địa chỉ người bán",
    "Người mua (đơn vị)",
    "MST người mua",
    "Hình thức TT",
    "Tiền hàng",
    "Thuế suất",
    "Tiền thuế",
    "Tổng tiền",
    "Trạng thái",
    "Ghi chú / Note",
]
LINE_COLUMNS = [
    "File",
    "Số HĐ",
    "STT",
    "Tên hàng hóa",
    "ĐVT",
    "Số lượng",
    "Đơn giá",
    "Thành tiền",
]
ISSUE_COLUMNS = ["File", "Số HĐ", "Lỗi (VI)", "Issue (EN)"]

STATUS_OK = "OK"
STATUS_ERROR = "LỖI / ERROR"

MONEY_COLUMNS = {
    "Tiền hàng",
    "Tiền thuế",
    "Tổng tiền",
    "Đơn giá",
    "Thành tiền",
}
TEXT_COLUMNS = {
    "Mẫu số",
    "Ký hiệu",
    "Số HĐ",
    "MST người bán",
    "MST người mua",
}

HEADER_FILL = PatternFill(start_color="1B3A4B", end_color="1B3A4B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
ERROR_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
ERROR_FONT = Font(name="Calibri", color="9C0006")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
TAB_COLORS = {"Invoices": "1B3A4B", "Line Items": "0F6E6B", "Issues": "9C0006"}


def build_parser():
    """Command-line arguments for the extractor."""
    parser = argparse.ArgumentParser(
        description="Extract supplier invoices from PDFs and write an Excel workbook."
    )
    parser.add_argument("input_dir", help="Folder that contains invoice PDF files")
    parser.add_argument(
        "-o",
        "--output",
        default=str(Path("output") / "invoices.xlsx"),
        help="Where to write the workbook (default: output/invoices.xlsx)",
    )
    return parser


def main(argv=None):
    """Read every PDF in the folder, validate it, and write the workbook."""
    configure_stdout()
    args = build_parser().parse_args(argv)
    input_dir = Path(args.input_dir)
    output_path = Path(args.output)

    if not input_dir.is_dir():
        print(f"Không thấy thư mục: {input_dir} / Input folder not found: {input_dir}")
        return 1

    pdf_paths = sorted(input_dir.glob("*.pdf"))
    if not pdf_paths:
        print(f"Không có file PDF trong {input_dir} / No PDF files in {input_dir}")
        return 1

    invoice_rows = []
    line_rows = []
    issue_rows = []

    for pdf_path in pdf_paths:
        invoice, issues = read_one(pdf_path)
        status = STATUS_OK if not issues else STATUS_ERROR
        invoice_rows.append(invoice_row(invoice, status, issues))
        line_rows.extend(line_item_rows(invoice))
        issue_rows.extend(issue_row(invoice, issue) for issue in issues)
        print_result(pdf_path.name, status, issues)

    try:
        write_workbook(invoice_rows, line_rows, issue_rows, output_path)
    except PermissionError:
        print(
            f"Không ghi được {output_path}. Hãy đóng file nếu đang mở trong Excel. / "
            f"Could not write {output_path}. Close the file if it is open in Excel."
        )
        return 1

    ok_count = sum(1 for row in invoice_rows if row["Trạng thái"] == STATUS_OK)
    error_count = len(invoice_rows) - ok_count
    print(
        f"Đã xử lý {len(invoice_rows)} hoá đơn: {ok_count} OK, {error_count} lỗi / "
        f"Processed {len(invoice_rows)} invoices: {ok_count} OK, {error_count} ERROR"
    )
    print(f"Đã ghi {output_path} / Wrote {output_path}")
    return 0


def configure_stdout():
    """Let the console print Vietnamese. The default Windows code page cannot."""
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except (OSError, ValueError):
            pass


def read_one(pdf_path):
    """Extract and validate one PDF. A broken file becomes an ERROR row."""
    try:
        invoice = extract_invoice(pdf_path)
    except Exception as exc:  # a bad PDF should not stop the rest of the folder
        invoice = {"file_name": pdf_path.name, "line_items": []}
        return invoice, [{"vi": f"Không đọc được PDF: {exc}", "en": f"Could not read PDF: {exc}"}]

    return invoice, validate_invoice(invoice)


def invoice_row(invoice, status, issues):
    """One Invoices-sheet row. The note is the Vietnamese summary, blank when OK."""
    return {
        "File": invoice.get("file_name"),
        "Mẫu số": empty_as_none(invoice.get("form_number")),
        "Ký hiệu": empty_as_none(invoice.get("serial")),
        "Số HĐ": empty_as_none(invoice.get("invoice_number")),
        "Ngày": parse_issue_date(invoice.get("issue_date")),
        "Người bán": invoice.get("seller_name"),
        "MST người bán": empty_as_none(invoice.get("seller_tax_code")),
        "Địa chỉ người bán": invoice.get("seller_address"),
        "Người mua (đơn vị)": invoice.get("buyer_unit"),
        "MST người mua": empty_as_none(invoice.get("buyer_tax_code")),
        "Hình thức TT": invoice.get("payment_method"),
        "Tiền hàng": invoice.get("subtotal"),
        "Thuế suất": invoice.get("vat_rate"),
        "Tiền thuế": invoice.get("vat_amount"),
        "Tổng tiền": invoice.get("total"),
        "Trạng thái": status,
        "Ghi chú / Note": note_text(issues),
    }


def note_text(issues):
    """Join the Vietnamese messages for the Ghi chú / Note column."""
    if not issues:
        return None
    return "; ".join(item["vi"] for item in issues)


def line_item_rows(invoice):
    """One Line Items row per extracted item."""
    rows = []
    for index, item in enumerate(invoice.get("line_items") or [], start=1):
        rows.append(
            {
                "File": invoice.get("file_name"),
                "Số HĐ": empty_as_none(invoice.get("invoice_number")),
                "STT": item.get("line_no") or index,
                "Tên hàng hóa": item.get("description"),
                "ĐVT": item.get("unit"),
                "Số lượng": item.get("qty"),
                "Đơn giá": item.get("unit_price"),
                "Thành tiền": item.get("amount"),
            }
        )
    return rows


def issue_row(invoice, issue):
    """One Issues-sheet row, with the Vietnamese and English wording side by side."""
    return {
        "File": invoice.get("file_name"),
        "Số HĐ": empty_as_none(invoice.get("invoice_number")),
        "Lỗi (VI)": issue["vi"],
        "Issue (EN)": issue["en"],
    }


def parse_issue_date(value):
    """Turn 12/01/2026 or 18.02.2026 into a real date. Unknown text is kept as-is."""
    if not value:
        return None
    for fmt in ("%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(str(value), fmt).date()
        except ValueError:
            continue
    return value


def empty_as_none(value):
    """Store a blank tax code as an empty Excel cell rather than the word None."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def write_workbook(invoice_rows, line_rows, issue_rows, output_path):
    """Write the three sheets, then apply number formats, widths, and ERROR highlighting."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    frames = {
        "Invoices": pd.DataFrame(invoice_rows, columns=INVOICE_COLUMNS),
        "Line Items": pd.DataFrame(line_rows, columns=LINE_COLUMNS),
        "Issues": pd.DataFrame(issue_rows, columns=ISSUE_COLUMNS),
    }
    # Tax codes and invoice numbers are text, so a leading zero is kept.
    for column in TEXT_COLUMNS:
        if column in frames["Invoices"].columns:
            frames["Invoices"][column] = frames["Invoices"][column].astype("string")
    frames["Line Items"]["Số HĐ"] = frames["Line Items"]["Số HĐ"].astype("string")
    frames["Issues"]["Số HĐ"] = frames["Issues"]["Số HĐ"].astype("string")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, frame in frames.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

    style_workbook(output_path)


def style_workbook(output_path):
    """Format the file that pandas just wrote."""
    from openpyxl import load_workbook

    workbook = load_workbook(output_path)
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        sheet.sheet_properties.tabColor = TAB_COLORS.get(sheet_name, "1B3A4B")
        style_sheet(sheet)
        if sheet_name == "Invoices":
            highlight_error_rows(sheet)
    workbook.save(output_path)


def style_sheet(sheet):
    """Bold header, borders, number formats, and column widths."""
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 22

    headers = [cell.value for cell in sheet[1]]
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=sheet.max_column):
        for cell, header in zip(row, headers):
            cell.border = THIN_BORDER
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="center", horizontal=column_alignment(header))
            apply_number_format(cell, header)

    autosize_columns(sheet, headers)


def column_alignment(header):
    """Right-align numbers, center short codes, left-align text."""
    if header in MONEY_COLUMNS or header in ("Số lượng", "STT", "Thuế suất"):
        return "right"
    if header in ("Trạng thái", "Ngày", "Ký hiệu", "Hình thức TT"):
        return "center"
    return "left"


def apply_number_format(cell, header):
    """VND amounts use #,##0. Dates and the VAT rate get their own formats."""
    if header in TEXT_COLUMNS or header == "Số HĐ":
        cell.number_format = "@"
    if cell.value is None:
        return
    if header in MONEY_COLUMNS or header == "Số lượng":
        cell.number_format = "#,##0"
    elif header == "Thuế suất":
        cell.number_format = '0"%"'
    elif header == "Ngày":
        cell.number_format = "DD/MM/YYYY"


def highlight_error_rows(sheet):
    """Fill every cell in an ERROR invoice row."""
    headers = [cell.value for cell in sheet[1]]
    try:
        status_index = headers.index("Trạng thái")
    except ValueError:
        return

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=sheet.max_column):
        if row[status_index].value != STATUS_ERROR:
            continue
        for cell in row:
            cell.fill = ERROR_FILL
            cell.font = ERROR_FONT


def autosize_columns(sheet, headers):
    """Set each column just wide enough for its contents, with a sensible cap."""
    caps = {
        "Lỗi (VI)": 88,
        "Issue (EN)": 84,
        "Ghi chú / Note": 88,
        "Tên hàng hóa": 36,
        "Người bán": 46,
        "Người mua (đơn vị)": 42,
        "Địa chỉ người bán": 56,
    }
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        longest = len(str(header))
        for cell in sheet[letter]:
            longest = max(longest, display_length(cell.value, header))
        cap = caps.get(header, 38)
        # Extra room so the filter arrow does not cover the header text.
        width = min(max(longest + 5, 14), cap)
        if header in MONEY_COLUMNS:
            width = max(width, 18)
        sheet.column_dimensions[letter].width = width


def display_length(value, header):
    """Length used for column width. Dates are measured as DD/MM/YYYY, not as timestamps."""
    if value is None:
        return 0
    if header == "Ngày" and hasattr(value, "strftime"):
        return len(value.strftime("%d/%m/%Y"))
    return len(str(value))


def print_result(file_name, status, issues):
    """Print one line per invoice, plus the problems in Vietnamese and English."""
    print(f"{file_name:<24} {status}")
    for item in issues:
        print(f"  - {item['vi']} / {item['en']}")


if __name__ == "__main__":
    sys.exit(main())
